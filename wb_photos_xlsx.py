#!/usr/bin/env python3
"""Excel-файл «Артикул продавца — Артикул WB — Фото» по всем карточкам Wildberries.

В столбец A пишется артикул продавца, в B — артикул WB, в C вставляются первые 3 фото карточки.
Токен берётся из переменной окружения WB_API_TOKEN (категория «Контент»).
Запуск: python3 wb_photos_xlsx.py [файл.xlsx]
"""
import io
import json
import os
import sys
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image

URL = "https://content-api.wildberries.ru/content/v2/get/cards/list"
PHOTOS = 3
W, H, GAP = 90, 120, 6  # размер фото в ячейке и отступ между ними, px


def fetch_cards(token):
    cards, cursor = [], {"limit": 100}
    while True:
        body = json.dumps({"settings": {"cursor": cursor, "filter": {"withPhoto": -1}}}).encode()
        req = urllib.request.Request(URL, data=body, method="POST", headers={
            "Authorization": token, "Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=60) as r:
            data = json.load(r)
        cards += data.get("cards", [])
        c = data.get("cursor", {})
        if c.get("total", 0) < cursor["limit"]:
            return cards
        cursor = {"limit": 100, "updatedAt": c["updatedAt"], "nmID": c["nmID"]}


def download(url, attempts=4):
    for n in range(attempts):
        try:
            with urllib.request.urlopen(url, timeout=60) as r:
                img = Image.open(io.BytesIO(r.read())).convert("RGB")
            break
        except Exception as e:
            if n == attempts - 1:
                print(f"Не удалось скачать {url}: {e}", file=sys.stderr)
                return None
            time.sleep(2 ** n)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return buf


def main():
    token = os.environ.get("WB_API_TOKEN")
    if not token:
        sys.exit("Не задана переменная окружения WB_API_TOKEN")
    out = sys.argv[1] if len(sys.argv) > 1 else "wb_photos.xlsx"

    cards = sorted(fetch_cards(token), key=lambda c: c.get("vendorCode", ""))
    urls = [[p["c246x328"] for p in c.get("photos", [])[:PHOTOS]] for c in cards]
    with ThreadPoolExecutor(16) as pool:
        images = list(pool.map(download, [u for row in urls for u in row]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Фото"
    ws.append(["Артикул продавца", "Артикул WB", "Фото"])
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = (PHOTOS * (W + GAP) + GAP) / 7
    ws.freeze_panes = "A2"

    it = iter(images)
    for i, (card, row_urls) in enumerate(zip(cards, urls), start=2):
        for col, value in ((1, card.get("vendorCode", "")), (2, card.get("nmID"))):
            cell = ws.cell(i, col, value)
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="center", horizontal="left")
        ws.row_dimensions[i].height = (H + 2 * GAP) * 0.75
        for j in range(len(row_urls)):
            buf = next(it)
            if buf is None:
                continue
            img = XLImage(buf)
            img.width, img.height = W, H
            marker = AnchorMarker(col=2, colOff=pixels_to_EMU(GAP + j * (W + GAP)),
                                  row=i - 1, rowOff=pixels_to_EMU(GAP))
            img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(
                pixels_to_EMU(W), pixels_to_EMU(H)))
            ws.add_image(img)

    wb.save(out)
    print(f"Готово: {out} — {len(cards)} артикулов")


if __name__ == "__main__":
    main()
