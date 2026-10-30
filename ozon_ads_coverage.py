#!/usr/bin/env python3
"""Ozon: склейки товаров, рекламируемые товары и склейки без активной рекламы.

Seller API (обязательно):      OZON_CLIENT_ID, OZON_API_KEY
Performance API (для списка товаров в рекламных кампаниях):
                               OZON_PERF_CLIENT_ID, OZON_PERF_CLIENT_SECRET
Без ключей Performance API отчёт строится частично: склейки и расход по кампаниям
(из начислений Seller API), но без привязки кампаний к товарам.

Запуск: python3 ozon_ads_coverage.py [out.xlsx] [дней=30]
"""
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SELLER = "https://api-seller.ozon.ru"
PERF = "https://api-performance.ozon.ru"

# Типы начислений (/v1/finance/accrual/types), относящиеся к рекламе
AD_ACCRUALS = {41: "Оплата за клик", 54: "Продвижение товара (оплата за заказ)", 75: "Трафареты",
               33: "Рекламные услуги", 3: "Продвижение бренда", 5: "Брендовая полка",
               19: "Внешнее продвижение", 23: "Реклама в сети Интернет", 87: "Реклама в соцсетях",
               130: "Медийная реклама", 49: "Продвижение Premium"}


def request(url, body=None, headers=None, method=None):
    data = json.dumps(body).encode() if body is not None else None
    for attempt in range(5):
        req = urllib.request.Request(url, data=data, method=method or ("POST" if data else "GET"),
                                     headers={"Content-Type": "application/json", **(headers or {})})
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < 4:
                time.sleep(2 ** attempt)
                continue
            sys.exit(f"{url}: HTTP {e.code} {e.read().decode()[:300]}")
        except (urllib.error.URLError, ConnectionError, TimeoutError):
            if attempt == 4:
                raise
            time.sleep(2 ** attempt)


def seller(path, body, creds):
    return request(SELLER + path, body, {"Client-Id": creds[0], "Api-Key": creds[1]})


def fetch_products(creds):
    ids, last = [], ""
    while True:
        r = seller("/v3/product/list", {"filter": {"visibility": "ALL"}, "limit": 1000, "last_id": last}, creds)["result"]
        ids += [x["product_id"] for x in r["items"]]
        if not r["items"] or len(ids) >= r["total"]:
            break
        last = r["last_id"]
    items = []
    for i in range(0, len(ids), 1000):
        items += seller("/v3/product/info/list", {"product_id": ids[i:i + 1000]}, creds)["items"]
    return items


def fetch_ad_spend(creds, days):
    """Расход по рекламным кампаниям из начислений: unit_number — ID кампании."""
    spend = defaultdict(lambda: {"type": "", "sum": 0.0, "days": set()})
    for k in range(1, days + 1):
        day = (date.today() - timedelta(days=k)).isoformat()
        last = None
        while True:
            body = {"date": day, **({"last_id": last} if last else {})}
            r = seller("/v1/finance/accrual/by-day", body, creds)
            for a in r["accruals"]:
                fee = a.get("non_item_fee")
                if fee and fee["type_id"] in AD_ACCRUALS:
                    s = spend[a.get("unit_number") or "—"]
                    s["type"] = AD_ACCRUALS[fee["type_id"]]
                    s["sum"] -= float(fee["accrued"]["amount"])
                    s["days"].add(a["date"])
            last = r.get("last_id")
            if not last or len(r["accruals"]) < 1000:
                break
    return spend


def fetch_perf_campaigns(perf_id, perf_secret):
    """Активные кампании и SKU в них. Возвращает ({id: кампания}, {sku: [id кампаний]})."""
    token = request(PERF + "/api/client/token", {"client_id": perf_id, "client_secret": perf_secret,
                                                  "grant_type": "client_credentials"})["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    campaigns = {str(c["id"]): c for c in request(PERF + "/api/client/campaign?state=CAMPAIGN_STATE_RUNNING", headers=h)["list"]}
    sku_campaigns = defaultdict(list)
    for cid, c in campaigns.items():
        if c.get("advObjectType") == "SEARCH_PROMO":
            continue
        page = 1
        while True:
            r = request(PERF + f"/api/client/campaign/{cid}/v2/products?page={page}&pageSize=100", headers=h)
            for p in r.get("products") or []:
                sku_campaigns[str(p["sku"])].append(cid)
            if len(r.get("products") or []) < 100:
                break
            page += 1
    # Оплата за заказ: товары не привязаны к отдельной кампании
    promo = [c for c in campaigns.values() if c.get("advObjectType") == "SEARCH_PROMO"]
    if promo:
        cid, page = str(promo[0]["id"]), 1
        while True:
            r = request(PERF + "/api/client/campaign/search_promo/v2/products", {"page": page, "pageSize": 100}, h)
            for p in r.get("products") or []:
                sku_campaigns[str(p["sku"])].append(cid)
            if len(r.get("products") or []) < 100:
                break
            page += 1
    return campaigns, sku_campaigns


def product_skus(p):
    return {str(p["sku"])} | {str(s["sku"]) for s in p.get("sources", [])}


def natural_key(s):
    """Сортировка артикулов с учётом чисел: «Агни 42» раньше «Агни 100»."""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def write_sheet(ws, header, rows, widths):
    ws.append(header)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    creds = (os.environ.get("OZON_CLIENT_ID"), os.environ.get("OZON_API_KEY"))
    if not all(creds):
        sys.exit("Не заданы OZON_CLIENT_ID / OZON_API_KEY")
    out = sys.argv[1] if len(sys.argv) > 1 else "ozon_ads_coverage.xlsx"
    days = int(sys.argv[2]) if len(sys.argv) > 2 else 30

    products = fetch_products(creds)
    spend = fetch_ad_spend(creds, days)
    perf = (os.environ.get("OZON_PERF_CLIENT_ID"), os.environ.get("OZON_PERF_CLIENT_SECRET"))
    campaigns, sku_campaigns = fetch_perf_campaigns(*perf) if all(perf) else ({}, None)

    models = defaultdict(list)
    for p in products:
        models[p["model_info"]["model_id"]].append(p)
    model_name = {m: Counter(p["name"] for p in ps).most_common(1)[0][0] for m, ps in models.items()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Склейки — артикулы"
    rows = []
    for m, ps in sorted(models.items(), key=lambda kv: -len(kv[1])):
        offers = sorted((p["offer_id"] for p in ps), key=natural_key)
        if sku_campaigns is None:
            ads = "нет данных (нужен доступ к Performance API)"
        else:
            advertised = sorted((p["offer_id"] for p in ps
                                 if any(sku_campaigns.get(s) for s in product_skus(p))), key=natural_key)
            ads = (f"Есть реклама ({len(advertised)} из {len(offers)}): " + ", ".join(advertised)
                   if advertised else "НЕТ АКТИВНОЙ РЕКЛАМЫ")
        rows.append([", ".join(offers), ads])
    write_sheet(ws, ["Склейка — артикулы", "Реклама на товары склейки"], rows, [110, 60])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Товары")
    rows = []
    for m, ps in sorted(models.items(), key=lambda kv: -len(kv[1])):
        for p in sorted(ps, key=lambda p: p["offer_id"]):
            skus = product_skus(p)
            stock = Counter()
            for s in p.get("stocks", {}).get("stocks", []):
                stock[s["source"]] += s.get("present", 0) - s.get("reserved", 0)
            if sku_campaigns is None:
                in_ads, cids = "н/д", ""
            else:
                cids = sorted({c for s in skus for c in sku_campaigns.get(s, [])})
                in_ads, cids = ("Да" if cids else "Нет"), ", ".join(cids)
            rows.append([m, model_name[m], p["offer_id"].rsplit(" ", 1)[0], p["offer_id"], p["sku"], p["name"],
                         float(p["price"]), stock["fbo"], stock["fbs"], in_ads, cids])
    write_sheet(ws, ["ID склейки", "Склейка", "Модель / цвет", "Артикул", "SKU", "Название", "Цена, ₽",
                     "Остаток FBO", "Остаток FBS", "В активной рекламе", "Кампании"],
                rows, [13, 40, 24, 26, 13, 45, 10, 11, 11, 12, 24])
    n = len(rows) + 1

    ws = wb.create_sheet("Склейки")
    rng = lambda col: f"Товары!${col}$2:${col}${n}"
    rows = []
    for i, (m, ps) in enumerate(sorted(models.items(), key=lambda kv: -len(kv[1])), 2):
        rows.append([m, model_name[m], f"=COUNTIF({rng('A')},A{i})", ps[0]["model_info"]["count"],
                     f"=SUMIF({rng('A')},A{i},{rng('H')})+SUMIF({rng('A')},A{i},{rng('I')})",
                     f'=COUNTIFS({rng("A")},A{i},{rng("J")},"Да")',
                     f'=IF(COUNTIFS({rng("A")},A{i},{rng("J")},"н/д")>0,"нет данных",'
                     f'IF(F{i}=0,"НЕТ АКТИВНОЙ РЕКЛАМЫ","есть реклама"))'])
    write_sheet(ws, ["ID склейки", "Склейка", "Товаров (активных)", "Товаров в склейке всего (вкл. архив)",
                     "Остаток FBO+FBS", "Товаров в рекламе", "Статус рекламы"], rows, [13, 45, 12, 14, 12, 12, 24])

    ws = wb.create_sheet("Рекламные кампании")
    rows = []
    for cid, s in sorted(spend.items(), key=lambda kv: -kv[1]["sum"]):
        c = campaigns.get(cid, {})
        rows.append([cid, c.get("title", ""), s["type"], c.get("advObjectType", ""), round(s["sum"], 2),
                     len(s["days"]), max(s["days"]),
                     sum(cid in v for v in sku_campaigns.values()) if sku_campaigns is not None else "н/д"])
    write_sheet(ws, ["ID кампании", "Название", "Тип списания", "Тип объекта", f"Расход за {days} дн., ₽",
                     "Дней со списаниями", "Последнее списание", "SKU в кампании"],
                rows, [13, 35, 34, 14, 14, 12, 14, 12])
    total = len(rows) + 2
    ws.cell(total, 1, "Итого").font = Font(name="Arial", bold=True)
    ws.cell(total, 5, f"=SUM(E2:E{total - 1})").font = Font(name="Arial", bold=True)
    for r in range(2, total + 1):
        ws.cell(r, 5).number_format = "#,##0"

    wb.save(out)
    print(f"Товаров: {len(products)}, склеек: {len(models)}, рекламных кампаний со списаниями: {len(spend)}")
    print(f"Расход на рекламу за {days} дн.: {sum(s['sum'] for s in spend.values()):,.0f} ₽".replace(",", " "))
    if sku_campaigns is None:
        print("Ключи Performance API не заданы — привязка рекламы к товарам не построена.")
    print(f"Файл: {out}")


if __name__ == "__main__":
    main()
