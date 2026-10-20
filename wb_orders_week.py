#!/usr/bin/env python3
"""Сводка заказов Wildberries за неделю по номенклатуре (размеры свёрнуты).

Одна строка на карточку (nmId = артикул продавца + цвет), все размеры суммируются.
Токен — из переменной окружения WB_API_TOKEN (категория «Статистика», только чтение).
Запуск: python3 wb_orders_week.py [ПОНЕДЕЛЬНИК YYYY-MM-DD] [-o файл.xlsx]
По умолчанию — прошлая неделя (пн–вс).
"""
import argparse
import json
import os
import sys
import urllib.request
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

URL = "https://statistics-api.wildberries.ru/api/v1/supplier/orders?dateFrom={d}&flag=0"


def fetch(since, token):
    req = urllib.request.Request(URL.format(d=since), headers={"Authorization": token})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.load(r)


def aggregate(orders):
    rows = defaultdict(lambda: {"art": "", "brand": "", "subject": "", "sizes": set(),
                                "n": 0, "cancel": 0, "sum": 0.0, "paid": 0.0})
    for o in orders:
        r = rows[o.get("nmId")]
        r["art"] = o.get("supplierArticle") or r["art"]
        r["brand"] = o.get("brand") or r["brand"]
        r["subject"] = o.get("subject") or r["subject"]
        r["sizes"].add(o.get("techSize") or "—")
        r["n"] += 1
        if o.get("isCancel"):
            r["cancel"] += 1
        r["sum"] += o.get("priceWithDisc", 0)
        r["paid"] += o.get("finishedPrice", 0)
    return sorted(rows.items(), key=lambda x: (-x[1]["n"], -x[1]["sum"]))


def write_xlsx(path, title, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы по номенклатуре"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    head = ["Артикул продавца", "Артикул WB (nmId)", "Бренд", "Предмет", "Размеров в заказах",
            "Заказано, шт", "Отменено, шт", "Без отмен, шт",
            "Сумма заказов, ₽ (цена со скидкой продавца)", "Оплачено покупателями, ₽"]
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for nm, r in rows:
        ws.append([r["art"], nm, r["brand"], r["subject"], len(r["sizes"]),
                   r["n"], r["cancel"], r["n"] - r["cancel"], round(r["sum"], 2), round(r["paid"], 2)])
    last = ws.max_row
    ws.append(["Итого", "", "", "", ""] + [f"=SUM({get_column_letter(c)}3:{get_column_letter(c)}{last})"
                                           for c in range(6, 11)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col in ("I", "J"):
        for c in ws[col][2:]:
            c.number_format = "# ##0"
    widths = [22, 14, 14, 22, 10, 12, 12, 12, 22, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    wb.save(path)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("monday", nargs="?")
    p.add_argument("-o", "--out")
    a = p.parse_args()
    token = os.environ.get("WB_API_TOKEN")
    if not token:
        sys.exit("Не задана переменная окружения WB_API_TOKEN")
    if a.monday:
        start = date.fromisoformat(a.monday)
    else:
        today = date.today()
        start = today - timedelta(days=today.weekday() + 7)
    end = start + timedelta(days=6)
    days = {(start + timedelta(i)).isoformat() for i in range(7)}

    seen, orders = set(), []
    for o in fetch(start.isoformat(), token):
        if o.get("date", "")[:10] in days and o.get("srid") not in seen:
            seen.add(o.get("srid"))
            orders.append(o)
    rows = aggregate(orders)
    period = f"{start:%d.%m}–{end:%d.%m.%Y}"
    out = a.out or f"wb_orders_{start}_{end}.xlsx"
    write_xlsx(out, f"Заказы WB за {period} (размеры свёрнуты)", rows)

    n = sum(r["n"] for _, r in rows)
    c = sum(r["cancel"] for _, r in rows)
    s = sum(r["sum"] for _, r in rows)
    print(f"Заказы WB за {period}: {n} шт (отмен {c}), сумма {s:,.0f} ₽, карточек {len(rows)}".replace(",", " "))
    for nm, r in rows[:10]:
        print(f"  {r['art']:<25} {r['n']:>5} шт  {r['sum']:>12,.0f} ₽".replace(",", " "))
    print(f"Файл: {out}")


if __name__ == "__main__":
    main()
