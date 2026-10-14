#!/usr/bin/env python3
"""Рентабельность Ozon по «Отчёту по начислениям».

Правила расчёта — docs/rentabelnost_wb_ozon.md, разделы 1–3, 5–7.

Запуск:
  python3 ozon_profitability.py НАЧИСЛЕНИЯ.xlsx СЕБЕСЫ.xlsx [-p ПРОШЛАЯ_НЕДЕЛЯ.xlsx]
                                [-o РЕЗУЛЬТАТ.xlsx] [--tax 0.10] [--credit-body 0]

Результат — Excel: лист «Рента Ozon» (прибыль, деньги на счёт, показатели,
комментарий), лист «Маржа по артикулам» (одна строка на модель) и, если есть,
лист «Нет себестоимости». Короткая сводка печатается в консоль.
"""
import argparse
import re
import sys
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Ручные соответствия модель → себестоимость (раздел 1 инструкции)
MANUAL_COST = {
    "Агни розовый варенка": 454.00,
    "Багги/ИВГПУтемно-синие": 910.00,
    "Карго Черный": 686.24,
    "Трикси2/черный": 863.60,
    "Фитнес голубой": 241.00,
}

# Статьи в порядке вывода
ITEMS = ["sales", "returns", "commission", "logistics", "ads", "other", "interest"]
ITEM_NAMES = {
    "sales": "  в т.ч. продажи",
    "returns": "  в т.ч. возвраты",
    "commission": "Комиссия Ozon",
    "logistics": "Логистика и услуги партнёров",
    "ads": "Реклама",
    "other": "Хранение, FBO, штрафы",
    "interest": "Проценты по кредиту",
}
GROUPS = {
    "продажи": "sales",
    "возвраты": "returns",
    "вознаграждение ozon": "commission",
    "услуги доставки": "logistics",
    "услуги партнеров": "logistics",
    "продвижение и реклама": "ads",
    "услуги fbo": "other",
    "другие услуги и штрафы": "other",
    "компенсации и декомпенсации": "other",
}
INTEREST_TYPES = ("досрочная выплата", "гибкий график выплат")
COLS = {
    "group": "группа услуг",
    "type": "тип начисления",
    "article": "артикул",
    "qty": "количество",
    "amount": "сумма итого",
}
SIZE_RE = re.compile(
    r"[\s_]+(\d{2,3}(?:[-/]\d{2,3})*|[2-6]?X{0,3}[SML]|ONE\s?SIZE|OS)$", re.I)


def rub(x):
    return f"{x:,.0f} ₽".replace(",", "\u00a0")


def pct(x, d=1):
    return f"{x:.{d}%}".replace(".", ",")


def pp(x):
    return f"{x * 100:+.1f}".replace(".", ",").replace("-", "−") + " п.п."


def norm(s):
    return re.sub(r"\s+", " ", str(s).replace("ё", "е").replace("Ё", "Е")).strip().lower()


def find_col(df, key, required=True):
    for c in df.columns:
        if norm(c).startswith(key):
            return c
    if required:
        sys.exit(f"В выгрузке нет колонки «{key}». Колонки: {list(df.columns)}")
    return None


def to_num(s):
    if s.dtype == object:
        s = s.astype(str).str.replace(" ", "").str.replace(" ", "").str.replace(",", ".")
    return pd.to_numeric(s, errors="coerce").fillna(0.0)


def model_of(article):
    return SIZE_RE.sub("", str(article).strip()).strip()


# ---------- себестоимость ----------

def load_costs(path):
    """Себестоимость: ключи «art:<артикул>» (точное совпадение) и «<модель>».

    Понимает два формата: «Цены_ВБ_себесы.xlsx» (лист «Лист_1», модель в скобках
    в «Номенклатуре», цена в «ВИД ЦЕНЫ (СТОП-ЦЕНА)») и прайс Ozon (колонки
    «Артикул» и «Себестоимость, руб» в шапке первой строки).
    """
    xl = pd.ExcelFile(path)
    prices = defaultdict(list)
    costs = {}
    if "Лист_1" in xl.sheet_names:
        df = pd.read_excel(path, sheet_name="Лист_1", header=1)
        name_col = find_col(df, "номенклатура")
        price_col = find_col(df, "вид цены")
        for name, price in zip(df[name_col], to_num(df[price_col])):
            m = re.findall(r"\(([^()]+)\)", str(name))
            if m and price > 0:
                prices[norm(m[-1])].append(float(price))
    else:
        df = pd.read_excel(path, sheet_name=xl.sheet_names[0])
        art_col = find_col(df, "артикул")
        price_col = find_col(df, "себестоимость")
        for art, price in zip(df[art_col], to_num(df[price_col])):
            if pd.notna(art) and price > 0:
                costs["art:" + norm(art)] = float(price)
                prices[norm(model_of(art))].append(float(price))
    conflicts = []
    for key, vals in prices.items():
        costs[key] = max(set(vals), key=vals.count)
        if len(set(vals)) > 1:
            conflicts.append((key, sorted(set(vals))))
    for model, price in MANUAL_COST.items():
        costs.setdefault(norm(model), price)
    return costs, conflicts


# ---------- отчёт по начислениям ----------

def load_report(path):
    xl = pd.ExcelFile(path)
    notes = []
    sheet = "Начисления"
    if sheet not in xl.sheet_names:
        sheet = xl.sheet_names[0]
        notes.append(f"Листа «Начисления» нет, взят лист «{sheet}».")
    head = pd.read_excel(path, sheet_name=sheet, header=None, nrows=1)
    period = " ".join(str(v) for v in head.iloc[0] if pd.notna(v)).strip()
    df = pd.read_excel(path, sheet_name=sheet, header=1)
    cols = {k: find_col(df, v) for k, v in COLS.items()}
    df = df.rename(columns={v: k for k, v in cols.items()})[list(COLS)]
    df = df[df["group"].notna() | df["type"].notna()].copy()
    df["amount"] = to_num(df["amount"])
    df["qty"] = to_num(df["qty"])
    df["article"] = df["article"].where(df["article"].notna(), "").astype(str).str.strip()
    df["article"] = df["article"].replace({"nan": "", "-": ""})
    df["group"] = df["group"].fillna("").astype(str)
    df["type"] = df["type"].fillna("").astype(str)

    def item(row):
        t = norm(row["type"])
        if any(x in t for x in INTEREST_TYPES):
            return "interest"
        return GROUPS.get(norm(row["group"]), "other")

    df["item"] = df.apply(item, axis=1)
    unknown = sorted({g for g in df.loc[df["item"] == "other", "group"] if norm(g) not in GROUPS})
    loans = df[(norm_series(df["group"]).str.contains("займ")) & (df["item"] != "interest")]
    if unknown:
        notes.append("Группы вне списка инструкции (ушли в «Хранение, FBO, штрафы»): "
                     + ", ".join(f"«{g}»" for g in unknown))
    if len(loans):
        notes.append(f"В «Займах» есть строки, кроме процентов: {loans['amount'].sum():,.2f} ₽ "
                     "— проверь, не тело ли это кредита.")
    return df, period, notes


def norm_series(s):
    return s.map(norm)


def compute(df, costs, tax):
    totals = {k: float(df.loc[df["item"] == k, "amount"].sum()) for k in ITEMS}
    grand = float(df["amount"].sum())
    t = norm_series(df["type"])
    sold = df[t == "выручка"].groupby("article")["qty"].sum().abs()
    returned = df[t == "возврат выручки"].groupby("article")["qty"].sum().abs()

    # По артикулам: прямые строки + разнесение строк без артикула по доле продаж
    direct = df[df["article"] != ""].pivot_table(
        index="article", columns="item", values="amount", aggfunc="sum").fillna(0.0)
    for k in ITEMS:
        if k not in direct:
            direct[k] = 0.0
    pool = df[df["article"] == ""].groupby("item")["amount"].sum()
    sales_sum = direct["sales"].sum()
    share = direct["sales"] / sales_sum if sales_sum else direct["sales"] * 0
    sku = direct[ITEMS].copy()
    for k, v in pool.items():
        sku[k] = sku[k] + share * v
    sku["sold"] = sold.reindex(sku.index).fillna(0)
    sku["returned"] = returned.reindex(sku.index).fillna(0)
    sku["model"] = [model_of(a) for a in sku.index]
    sku["price"] = [costs.get("art:" + norm(a), costs.get(norm(m)))
                    for a, m in zip(sku.index, sku["model"])]
    missing = sku[sku["price"].isna() & ((sku["sold"] > 0) | (sku["returned"] > 0))]
    sku["cost"] = (sku["sold"] - sku["returned"]) * sku["price"].fillna(0)
    sku["revenue"] = sku["sales"] + sku["returns"]
    sku["margin"] = sku[ITEMS].sum(axis=1) - sku["cost"] - tax * sku["revenue"]

    revenue = totals["sales"] + totals["returns"]
    cost = float(sku["cost"].sum())
    margin = grand - cost - tax * revenue
    with_sales = sku[sku["sold"] > 0]
    res = dict(
        totals=totals, grand=grand, revenue=revenue, cost=cost, tax=tax * revenue,
        margin=margin, margin_nc=margin - totals["interest"],
        sold=float(sold.sum()), returned=float(returned.sum()),
        sku_sales=int(len(with_sales)), sku_minus=int((sku["margin"] < 0).sum()),
        loss=float(sku.loc[sku["margin"] < 0, "margin"].sum()),
        sku=sku, missing=missing, unallocated_check=float(sku[ITEMS].sum().sum()) - grand,
    )
    return res


# ---------- Excel ----------

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FBLUE = Font(name=FONT, size=10, color="0000FF")
FH = Font(name=FONT, size=12, bold=True)
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="E7E6E6")
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6EFCE")
THIN = Border(bottom=Side(style="thin", color="999999"))
RUB = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
PCT2 = '0.00%;-0.00%;"-"'


def put(ws, ref, value, font=F, fmt=None, fill=None, bold=False):
    c = ws[ref]
    c.value = value
    c.font = FB if bold else font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def write_main(ws, res, prev, period, tax, credit_body, comments, notes):
    ws.title = "Рента Ozon"
    put(ws, "A1", f"Рентабельность Ozon. {period}", FH)
    put(ws, "A2", "Ставка налога (от выручки)")
    put(ws, "B2", tax, FBLUE, PCT, YELLOW)
    ws["B2"].comment = Comment("По умолчанию 10% от выручки (продажи + возвраты). Меняй здесь.", "Claude")

    put(ws, "A4", "ПРИБЫЛЬ ЗА НЕДЕЛЮ", FB)
    for col, h in zip("BCD", ["Факт", "Без кредита", "Прошлая неделя"]):
        put(ws, f"{col}4", h, bold=True, fill=GREY)
    ws["A4"].fill = GREY
    r = 5
    rows = {}
    put(ws, f"A{r}", "Выручка (продажи + возвраты)", bold=True)
    rows["revenue"] = r
    r += 1
    for k in ITEMS:
        rows[k] = r
        put(ws, f"A{r}", ITEM_NAMES[k])
        r += 1
    rows["grand"], rows["cost"], rows["tax"] = r, r + 1, r + 2
    rows["margin"], rows["rent"], rows["roi"] = r + 3, r + 4, r + 5
    labels = {
        "grand": "Итого начислено", "cost": "Себестоимость", "tax": "Налог",
        "margin": "Маржа", "rent": "Рентабельность к выручке", "roi": "ROI к закупке",
    }
    for k, lbl in labels.items():
        put(ws, f"A{rows[k]}", lbl, bold=k in ("grand", "margin"))

    for col, data in (("B", res), ("C", res), ("D", prev)):
        if data is None:
            continue
        nc = col == "C"
        for k in ITEMS:
            v = 0 if (nc and k == "interest") else round(data["totals"][k], 2)
            if nc and k != "interest":
                put(ws, f"C{rows[k]}", f"=B{rows[k]}", fmt=RUB)
            else:
                put(ws, f"{col}{rows[k]}", v, FBLUE, RUB)
        rv, gi, ct = rows["revenue"], rows["grand"], rows["cost"]
        put(ws, f"{col}{rv}", f"={col}{rows['sales']}+{col}{rows['returns']}", fmt=RUB, bold=True)
        put(ws, f"{col}{gi}", f"=SUM({col}{rows['sales']}:{col}{rows['interest']})", fmt=RUB, bold=True)
        if nc:
            put(ws, f"C{ct}", f"=B{ct}", fmt=RUB)
        else:
            put(ws, f"{col}{ct}", -round(data["cost"], 2), FBLUE, RUB)
        put(ws, f"{col}{rows['tax']}", f"=-{col}{rv}*$B$2", fmt=RUB)
        put(ws, f"{col}{rows['margin']}", f"={col}{gi}+{col}{ct}+{col}{rows['tax']}", fmt=RUB, bold=True)
        put(ws, f"{col}{rows['rent']}", f"=IF({col}{rv}=0,0,{col}{rows['margin']}/{col}{rv})", fmt=PCT2)
        put(ws, f"{col}{rows['roi']}", f"=IF({col}{ct}=0,0,-{col}{rows['margin']}/{col}{ct})", fmt=PCT)
    ws[f"B{rows['cost']}"].comment = Comment(
        "(продано − возвращено) × стоп-цена модели из «Цены_ВБ_себесы.xlsx»", "Claude")
    ws[f"B{rows['interest']}"].comment = Comment(
        "Типы «Досрочная выплата» + «Гибкий график выплат». Тело кредита сюда не входит.", "Claude")
    for k in ("grand", "margin"):
        for col in "ABCD":
            ws[f"{col}{rows[k]}"].border = THIN
    r = rows["roi"] + 1
    put(ws, f"A{r}", "Сверка: итог выгрузки «Сумма итого»")
    put(ws, f"B{r}", round(res["grand"], 2), FBLUE, RUB)
    put(ws, f"C{r}", f'=IF(ABS(B{r}-B{rows["grand"]})<1,"сходится","НЕ СХОДИТСЯ")')

    # Деньги на счёт
    r += 2
    put(ws, f"A{r}", "ДЕНЬГИ НА СЧЁТ", FB, fill=GREY)
    put(ws, f"B{r}", "₽", bold=True, fill=GREY)
    put(ws, f"A{r+1}", "Начислено")
    put(ws, f"B{r+1}", f"=B{rows['grand']}", fmt=RUB)
    put(ws, f"A{r+2}", "Гашение тела кредита (из кабинета)")
    put(ws, f"B{r+2}", credit_body, FBLUE, RUB, YELLOW)
    ws[f"B{r+2}"].comment = Comment(
        "Ввести сумму «Займы и факторинг» из кабинета (положительным числом). "
        "В маржу не входит — это возврат аванса, не расход.", "Claude")
    put(ws, f"A{r+3}", "Поступило на счёт", bold=True)
    put(ws, f"B{r+3}", f"=B{r+1}-B{r+2}", fmt=RUB, bold=True)

    # Показатели
    r += 5
    put(ws, f"A{r}", "ПОКАЗАТЕЛИ", FB, fill=GREY)
    put(ws, f"B{r}", "Факт", bold=True, fill=GREY)
    put(ws, f"C{r}", "Прошлая неделя", bold=True, fill=GREY)
    rv = rows["revenue"]
    ind = [
        ("Продано, шт", "sold", "#,##0"),
        ("Возвраты, шт", "returned", "#,##0"),
        ("Возвраты, %", "=IF({c}{s}=0,0,{c}{t}/{c}{s})", PCT),
        ("Средний чек (продажи / шт)", "=IF({c}{s}=0,0,{m}{sales}/{c}{s})", RUB),
        ("Комиссия к выручке", "=IF({m}{rv}=0,0,-{m}{commission}/{m}{rv})", PCT),
        ("Логистика к выручке", "=IF({m}{rv}=0,0,-{m}{logistics}/{m}{rv})", PCT),
        ("ДРР (реклама к выручке)", "=IF({m}{rv}=0,0,-{m}{ads}/{m}{rv})", PCT),
        ("Себестоимость к выручке", "=IF({m}{rv}=0,0,-{m}{cost}/{m}{rv})", PCT),
        ("SKU с продажами", "sku_sales", "#,##0"),
        ("SKU в минусе", "sku_minus", "#,##0"),
        ("Убыток по SKU в минусе", "loss", RUB),
        ("Проценты по кредиту к марже", "=IF({m}{margin}=0,0,-{m}{interest}/{m}{margin})", PCT),
    ]
    first = r + 1
    for i, (lbl, src, fmt) in enumerate(ind):
        rr = first + i
        put(ws, f"A{rr}", lbl)
        for col, main_col, data in (("B", "B", res), ("C", "D", prev)):
            if data is None:
                continue
            if src.startswith("="):
                f = src.format(c=col, m=main_col, s=first, t=first + 1, rv=rv,
                               **{k: v for k, v in rows.items() if k != "rv"})
                put(ws, f"{col}{rr}", f, fmt=fmt)
            else:
                put(ws, f"{col}{rr}", round(data[src], 2), FBLUE, fmt)
    ws[f"B{first + 8}"].comment = Comment("SKU = артикул Ozon (с размером). Строки без артикула "
                                          "разнесены пропорционально продажам.", "Claude")

    # Комментарий справа
    put(ws, "F4", "КОММЕНТАРИЙ", FB, fill=GREY)
    for i, line in enumerate(comments + [f"⚠ {n}" for n in notes]):
        c = put(ws, f"F{5 + i}", f"• {line}")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 38
    for col in "BCD":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["E"].width = 3
    ws.column_dimensions["F"].width = 80
    ws.freeze_panes = "A5"


SKU_HEAD = ["Наименование", "Продано шт", "Продажа ₽", "Возврат шт", "Возврат ₽",
            "Комиссия", "Логистика", "Хранение, FBO, штрафы", "Реклама ₽", "Реклама %",
            "Проценты кредита", "Стоп-цена", "СТОП-ЦЕНА", "Итого", "Налог", "ИТОГОВЫЙ ИТОГ",
            "Маржа", "Рентабельность", "Маржа на ед.", "Статус"]


def write_models(ws, res, tax_ref):
    sku = res["sku"]
    g = sku.groupby("model")
    models = g[ITEMS + ["sold", "returned", "cost"]].sum()
    # средневзвешенная стоп-цена: у размеров одной модели цены могут различаться
    units = models["sold"] - models["returned"]
    models["price"] = (models["cost"] / units).where(units != 0, g["price"].first())
    models.loc[g["price"].apply(lambda s: s.isna().all()), "price"] = None
    models = models.sort_index(key=lambda s: s.map(norm))

    put(ws, "A1", "Маржа на артикул (модель, все размеры вместе)", FH)
    put(ws, "A2", "Ставка налога")
    put(ws, "B2", f"={tax_ref}", fmt=PCT)
    put(ws, "D2", "Синие — данные из выгрузки, жёлтые — ввод (стоп-цена), остальное — формулы. "
                  "Расходы со знаком плюс. Рентабельность = Маржа / СТОП-ЦЕНА (к закупке).", F)
    for i, h in enumerate(SKU_HEAD, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font, c.fill = FB, GREY
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 30
    r0 = 5
    for i, (model, d) in enumerate(models.iterrows()):
        r = r0 + i
        vals = [model, d["sold"], d["sales"], d["returned"], -d["returns"], -d["commission"],
                -d["logistics"], -d["other"], -d["ads"], None, -d["interest"],
                d["price"] if pd.notna(d["price"]) else None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F if j == 1 else FBLUE
            c.number_format = RUB if j > 1 else "General"
        ws.cell(row=r, column=12).fill = YELLOW
        net = f"(C{r}-E{r})"
        formulas = {
            10: (f"=IF({net}=0,0,I{r}/{net})", PCT),
            13: (f"=L{r}*(B{r}-D{r})", RUB),
            14: (f"={net}-F{r}-G{r}-H{r}-I{r}-K{r}", RUB),
            15: (f"={net}*$B$2", RUB),
            16: (f"=N{r}-O{r}", RUB),
            17: (f"=P{r}-M{r}", RUB),
            18: (f"=IF(M{r}=0,0,Q{r}/M{r})", PCT),
            19: (f"=IF((B{r}-D{r})=0,0,Q{r}/(B{r}-D{r}))", RUB),
            20: (f'=IF(AND(L{r}="",B{r}-D{r}<>0),"НЕТ СЕБЕСТОИМОСТИ",IF(Q{r}<0,"В МИНУСЕ","в плюсе"))', None),
        }
        for j, (f, fmt) in formulas.items():
            c = ws.cell(row=r, column=j, value=f)
            c.font = F
            if fmt:
                c.number_format = fmt
    last = r0 + len(models) - 1
    tr = last + 1
    put(ws, f"A{tr}", "ИТОГО", bold=True)
    for j in range(2, 20):
        col = ws.cell(row=4, column=j).column_letter
        if j in (10, 12, 18, 19):
            continue
        c = ws.cell(row=tr, column=j, value=f"=SUM({col}{r0}:{col}{last})")
        c.font, c.number_format = FB, RUB
    ws[f"J{tr}"] = f"=IF((C{tr}-E{tr})=0,0,I{tr}/(C{tr}-E{tr}))"
    ws[f"R{tr}"] = f"=IF(M{tr}=0,0,Q{tr}/M{tr})"
    ws[f"S{tr}"] = f"=IF((B{tr}-D{tr})=0,0,Q{tr}/(B{tr}-D{tr}))"
    for ref, fmt in ((f"J{tr}", PCT), (f"R{tr}", PCT), (f"S{tr}", RUB)):
        ws[ref].font, ws[ref].number_format = FB, fmt
    for col in range(1, 21):
        ws.cell(row=tr, column=col).border = Border(top=Side(style="thin"))
    rng = f"T{r0}:T{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($L{r0}="",$B{r0}-$D{r0}<>0)'], fill=YELLOW, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$Q{r0}<0'], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$Q{r0}>=0'], fill=GREEN))

    # Экономика одной единицы
    u = tr + 2
    put(ws, f"A{u}", "ЭКОНОМИКА ОДНОЙ ЕДИНИЦЫ (в среднем по выкупленным)", FB, fill=GREY)
    put(ws, f"B{u}", "₽ на ед.", bold=True, fill=GREY)
    put(ws, f"C{u}", "% к цене", bold=True, fill=GREY)
    lines = [("Цена (выручка)", f"(C{tr}-E{tr})"), ("Комиссия", f"F{tr}"), ("Логистика", f"G{tr}"),
             ("Хранение, FBO, штрафы", f"H{tr}"), ("Реклама", f"I{tr}"), ("Проценты кредита", f"K{tr}"),
             ("Налог", f"O{tr}"), ("Себестоимость", f"M{tr}"), ("Маржа", f"Q{tr}")]
    units = f"(B{tr}-D{tr})"
    for i, (lbl, expr) in enumerate(lines, 1):
        rr = u + i
        put(ws, f"A{rr}", lbl, bold=lbl == "Маржа")
        put(ws, f"B{rr}", f"=IF({units}=0,0,{expr}/{units})", fmt=RUB, bold=lbl == "Маржа")
        put(ws, f"C{rr}", f"=IF(B{u+1}=0,0,B{rr}/B{u+1})", fmt=PCT)
    ws.column_dimensions["A"].width = 34
    for j in range(2, 21):
        ws.column_dimensions[ws.cell(row=4, column=j).column_letter].width = 12
    ws.freeze_panes = "B5"


def write_missing(wb, res):
    m = res["missing"]
    if m.empty:
        return
    ws = wb.create_sheet("Нет себестоимости")
    put(ws, "A1", "Не нашлась себестоимость — маржа по этим товарам завышена", FH)
    for i, h in enumerate(["Артикул Ozon", "Модель (после обрезки размера)", "Продано шт",
                           "Возврат шт", "Выручка ₽"], 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font, c.fill = FB, GREY
    for i, (art, d) in enumerate(m.sort_index().iterrows(), 4):
        for j, v in enumerate([art, d["model"], d["sold"], d["returned"],
                               round(d["sales"] + d["returns"], 2)], 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F
            if j == 5:
                c.number_format = RUB
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 35
    for col in "CDE":
        ws.column_dimensions[col].width = 13


# ---------- выводы ----------

def ratios(d):
    rv = d["revenue"] or 1
    t = d["totals"]
    return dict(rent=d["margin"] / rv, roi=d["margin"] / d["cost"] if d["cost"] else 0,
                comm=-t["commission"] / rv, log=-t["logistics"] / rv, drr=-t["ads"] / rv,
                cost=d["cost"] / rv, ret=d["returned"] / d["sold"] if d["sold"] else 0)


def make_comments(res, prev):
    r = ratios(res)
    out = []
    if prev:
        p = ratios(prev)
        dv = (res["revenue"] / prev["revenue"] - 1) if prev["revenue"] else 0
        out.append(f"Выручка {rub(res['revenue'])} ({pct(dv).replace('-', '−') if dv < 0 else '+' + pct(dv)} "
                   f"к прошлой неделе), маржа {rub(res['margin'])} (было {rub(prev['margin'])}).")
        out.append(f"Рентабельность {pct(r['rent'], 2)} (было {pct(p['rent'], 2)}), "
                   f"ROI к закупке {pct(r['roi'])} (было {pct(p['roi'])}).")
        diffs = [("комиссия", "comm"), ("логистика", "log"), ("ДРР", "drr"),
                 ("себестоимость", "cost"), ("возвраты", "ret")]
        moved = sorted(diffs, key=lambda x: -abs(r[x[1]] - p[x[1]]))[:3]
        for name, k in moved:
            out.append(f"{name[0].upper() + name[1:]}: {pct(r[k])} против {pct(p[k])} "
                       f"({pp(r[k] - p[k])}).")
    else:
        out.append(f"Выручка {rub(res['revenue'])}, маржа {rub(res['margin'])}: "
                   f"{pct(r['rent'], 2)} к выручке, ROI {pct(r['roi'])} к закупке.")
        out.append(f"Комиссия {pct(r['comm'])}, логистика {pct(r['log'])}, ДРР {pct(r['drr'])}, "
                   f"себестоимость {pct(r['cost'])}, возвраты {pct(r['ret'])}.")
    out.append(f"SKU в минусе: {res['sku_minus']} из {len(res['sku'])}, убыток по ним {rub(res['loss'])}.")
    if res["margin"]:
        out.append(f"Проценты по кредиту {rub(-res['totals']['interest'])} "
                   f"({pct(-res['totals']['interest'] / res['revenue'])} выручки); без них маржа "
                   f"{rub(res['margin_nc'])}.")
    # здравый смысл (раздел 7 инструкции)
    bounds = [("комиссия", "comm", .45, .55), ("логистика", "log", .08, .18),
              ("ДРР", "drr", .01, .05), ("себестоимость", "cost", .14, .21), ("возвраты", "ret", .03, .10)]
    odd = [f"{n} {pct(r[k])}" for n, k, lo, hi in bounds if not lo <= r[k] <= hi]
    if odd:
        out.append("Вне привычных рамок: " + ", ".join(odd) + " — стоит проверить.")
    return out[:7]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("report")
    ap.add_argument("costs")
    ap.add_argument("-p", "--prev", help="отчёт по начислениям за прошлую неделю")
    ap.add_argument("-o", "--out", default="Рента_Ozon.xlsx")
    ap.add_argument("--tax", type=float, default=0.10)
    ap.add_argument("--credit-body", type=float, default=0.0,
                    help="гашение тела кредита из кабинета, ₽ (только для «Деньги на счёт»)")
    a = ap.parse_args()

    costs, conflicts = load_costs(a.costs)
    df, period, notes = load_report(a.report)
    res = compute(df, costs, a.tax)
    prev = None
    if a.prev:
        pdf, _, pnotes = load_report(a.prev)
        prev = compute(pdf, costs, a.tax)
        notes += [f"Прошлая неделя: {n}" for n in pnotes]
    if conflicts:
        notes.append("У моделей разные стоп-цены по размерам (взята частая): "
                     + "; ".join(f"{k}: {v}" for k, v in conflicts[:5]))
    if not res["missing"].empty:
        notes.append(f"Нет себестоимости у {len(res['missing'])} артикулов — см. лист «Нет себестоимости».")
    assert abs(res["unallocated_check"]) < 0.01, "сумма по артикулам не равна итогу выгрузки"
    assert abs(sum(res["totals"].values()) - res["grand"]) < 0.01, "статьи не равны итогу"

    wb = Workbook()
    comments = make_comments(res, prev)
    write_main(wb.active, res, prev, period, a.tax, a.credit_body, comments, notes)
    write_models(wb.create_sheet("Маржа по артикулам"), res, "'Рента Ozon'!$B$2")
    write_missing(wb, res)
    wb.calculation.fullCalcOnLoad = True
    wb.save(a.out)

    rr = ratios(res)
    t = res["totals"]
    print(f"# Рентабельность Ozon. {period}\n")
    lines = [("Выручка", res["revenue"]), ("Комиссия Ozon", t["commission"]),
             ("Логистика и партнёры", t["logistics"]), ("Реклама", t["ads"]),
             ("Хранение, FBO, штрафы", t["other"]), ("Проценты по кредиту", t["interest"]),
             ("Итого начислено", res["grand"]), ("Себестоимость", -res["cost"]),
             ("Налог", -res["tax"]), ("Маржа", res["margin"]), ("Маржа без кредита", res["margin_nc"])]
    for name, v in lines:
        print(f"{name:<26}{rub(v):>16}")
    print(f"{'Рентабельность к выручке':<26}{pct(rr['rent'], 2):>16}")
    print(f"{'ROI к закупке':<26}{pct(rr['roi']):>16}")
    print("\n" + "\n".join(f"• {c}" for c in comments))
    for n in notes:
        print(f"⚠ {n}")
    if not res["missing"].empty:
        print("\nНет себестоимости:")
        for art, d in res["missing"].sort_index().iterrows():
            print(f"  {art}  (модель «{d['model']}», продано {d['sold']:.0f})")
    print(f"\nФайл: {a.out}")


if __name__ == "__main__":
    main()
